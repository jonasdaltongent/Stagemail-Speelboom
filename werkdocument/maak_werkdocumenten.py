#!/usr/bin/env python3
"""
maak_werkdocumenten.py
Genereert het werkdocument voor les 3 "Professionele e-mail aan een stageplaats" (De Speelboom, klas 3MWb):

  Stagemail.docx    het werkdocument dat de leerling INLEVERT (wordt in Classroom een Google-document)

Gebruik:  python3 maak_werkdocumenten.py     (vereist python-docx)

Waarom geen zip en geen Gmail: de leerlingen hebben met hun schoolaccount geen Gmail (Jonas, 09-10-2026),
en per les gebruiken ze één werkplek naast Classroom en de lespagina (_afspraken/didactiek.md). De e-mail
komt daarom in dit document, in een venster dat eruitziet als "Nieuw bericht" in Gmail.

LET OP bij aanpassen:
 - Deel 1, rij 1 is het uitgewerkte voorbeeld: dat doet de leraar voor (dia 6). Laat het ingevuld.
 - De adressen eindigen op .example: dat domein bestaat nooit echt (RFC 2606). Geen echte organisatie.
 - Alle namen en gegevens zijn fictief.
"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = RGBColor(0x2A, 0x39, 0x73)
GREY = RGBColor(0x5F, 0x63, 0x68)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
HERE = os.path.dirname(os.path.abspath(__file__))

MENTOR = ("mevrouw Lien Verhoeven", "lien.verhoeven@despeelboom.example")
BEGELEIDER = ("mevrouw Ilse Wouters", "ilse.wouters@school.example")


# ---------- hulpfuncties ----------
def basis_document(marge=2.0):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Arial"
    st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(marge)
        s.left_margin = s.right_margin = Cm(marge)
    return doc


def kop(doc, tekst, grootte=16, ruimte_voor=14):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(ruimte_voor)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(tekst)
    r.bold = True
    r.font.size = Pt(grootte)
    r.font.color.rgb = NAVY
    return p


def tekst(doc, s, cursief=False, klein=False, vet=False, na=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(na)
    r = p.add_run(s)
    r.italic = cursief
    r.bold = vet
    if klein:
        r.font.size = Pt(9.5)
        r.font.color.rgb = GREY
    return p


def schaduw(cel, kleur="E4E8F6"):
    tcPr = cel._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), kleur)
    tcPr.append(shd)


def randen(tabel, buiten="9AA0A6", binnen=None):
    """Randen van een tabel: alleen een kader rondom, of ook lijnen binnenin."""
    tblPr = tabel._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for kant in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + kant)
        kleur = buiten if kant in ("top", "left", "bottom", "right") else binnen
        if kleur:
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), "8")
            e.set(qn("w:color"), kleur)
        else:
            e.set(qn("w:val"), "nil")
        b.append(e)
    tblPr.append(b)


def onderrand(cel, kleur="DADCE0"):
    tcPr = cel._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    e = OxmlElement("w:bottom")
    e.set(qn("w:val"), "single")
    e.set(qn("w:sz"), "6")
    e.set(qn("w:color"), kleur)
    b.append(e)
    tcPr.append(b)


def breedtes(tabel, cm):
    """Vaste kolombreedtes, zodat Google Documenten de tabel niet zelf herschikt."""
    tabel.autofit = False
    for kol, w in zip(tabel.columns, cm):
        kol.width = Cm(w)
    for rij in tabel.rows:
        for cel, w in zip(rij.cells, cm):
            cel.width = Cm(w)


def celtekst(cel, s, vet=False, kleur=None, grootte=None):
    cel.text = ""
    r = cel.paragraphs[0].add_run(s)
    r.bold = vet
    if kleur:
        r.font.color.rgb = kleur
    if grootte:
        r.font.size = Pt(grootte)
    return r


def hoogte(rij, cm):
    trPr = rij._tr.get_or_add_trPr()
    h = OxmlElement("w:trHeight")
    h.set(qn("w:val"), str(int(cm * 567)))
    h.set(qn("w:hRule"), "atLeast")
    trPr.append(h)


def antwoordlijnen(doc, aantal=2):
    for _ in range(aantal):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.add_run("_" * 70)


def venster(doc, delen):
    """Een venster dat eruitziet als 'Nieuw bericht' in Gmail.
    delen: de grijze labels links van de tekstvakken in het bericht."""
    rijen = 1 + 4 + len(delen) + 1
    t = doc.add_table(rows=rijen, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    randen(t, buiten="9AA0A6")
    breedtes(t, [3.2, 13.8])

    kopcel = t.rows[0].cells[0].merge(t.rows[0].cells[1])
    schaduw(kopcel, "F2F6FC")
    celtekst(kopcel, "Nieuw bericht", vet=True, kleur=RGBColor(0x20, 0x21, 0x24))

    for i, veld in enumerate(["Aan", "Cc", "Bcc", "Onderwerp"], start=1):
        links, rechts = t.rows[i].cells
        celtekst(links, veld, kleur=GREY)
        onderrand(links)
        onderrand(rechts)
        hoogte(t.rows[i], 0.8)

    for j, label in enumerate(delen):
        rij = t.rows[5 + j]
        celtekst(rij.cells[0], label, kleur=GREY, grootte=8.5)
        hoogte(rij, 1.3 if j else 0.9)

    voet = t.rows[-1]
    knop = voet.cells[0]
    schaduw(knop, "0B57D0")
    celtekst(knop, "Verzenden", vet=True, kleur=WHITE)
    knop.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    celtekst(voet.cells[1], "Niet echt: je levert dit document in, in Classroom.", kleur=GREY, grootte=9)
    return t


# ---------- het werkdocument ----------
def stagemail():
    doc = basis_document()

    p = doc.add_paragraph()
    r = p.add_run("STAGEMAIL")
    r.bold = True
    r.font.size = Pt(22)
    r.font.color.rgb = NAVY
    tekst(doc, "Les 3 — Professionele e-mail aan een stageplaats  ·  Toegepaste Informatica  ·  De Speelboom",
          klein=True, na=10)

    t = doc.add_table(rows=1, cols=3)
    t.style = "Table Grid"
    for i, label in enumerate(["Naam:", "Klas:", "Datum:"]):
        c = t.rows[0].cells[i]
        c.text = label + " "
        c.paragraphs[0].runs[0].bold = True
    breedtes(t, [7.0, 4.0, 6.0])
    doc.add_paragraph()

    kop(doc, "Zo werk je", 13, ruimte_voor=6)
    for s in [
        "1.  Op de lespagina lees je wat je moet doen, stap voor stap.",
        "2.  Je schrijft alles in DIT document. Je hebt niets anders nodig.",
        "3.  Je e-mail schrijf je in deel 3, in een venster dat eruitziet zoals in Gmail.",
        "4.  Klaar? Klik in de opdracht op Inleveren. Er wordt niets echt verstuurd.",
    ]:
        p = doc.add_paragraph(s)
        p.paragraph_format.space_after = Pt(2)

    # ---- Je stage (het scenario)
    kop(doc, "Je stage", 13)
    k = doc.add_table(rows=1, cols=1)
    randen(k, buiten="A7B1DE")
    breedtes(k, [17.0])
    cel = k.rows[0].cells[0]
    schaduw(cel, "EEF0F9")
    cel.text = ""
    regels = [
        ("Je loopt stage in buitenschoolse opvang De Speelboom, van maandag 16 november tot vrijdag "
         "4 december 2026. Je kent daar nog niemand.", False),
        ("Je mentor op De Speelboom:  " + MENTOR[0] + "  —  " + MENTOR[1], True),
        ("Je stagebegeleider van school:  " + BEGELEIDER[0] + "  —  " + BEGELEIDER[1], True),
        ("Je opdracht: stuur je mentor een e-mail om kennis te maken. Vraag wanneer je eens kan "
         "langskomen. Je stagebegeleider wil meelezen.", False),
        ("Alle namen en adressen zijn verzonnen.", None),
    ]
    for n, (s, vet) in enumerate(regels):
        pp = cel.paragraphs[0] if n == 0 else cel.add_paragraph()
        pp.paragraph_format.space_after = Pt(4)
        rr = pp.add_run(s)
        if vet:
            rr.bold = True
        if vet is None:
            rr.italic = True
            rr.font.size = Pt(9)
            rr.font.color.rgb = GREY

    # ---- Deel 1
    kop(doc, "Deel 1 — Wat is er mis met de e-mail van Kobe?")
    tekst(doc, "De e-mail van Kobe staat op de lespagina, in stap 2. Schrijf drie fouten op, en hoe het beter "
               "kan. De eerste deed je leraar voor.", klein=True)
    t = doc.add_table(rows=4, cols=3)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(["", "Wat is er mis?", "Zo is het beter"]):
        c = t.rows[0].cells[i]
        c.text = h
        if h:
            c.paragraphs[0].runs[0].bold = True
        schaduw(c)
    rijen = [
        ("1", "Er is geen onderwerp.", "Kennismaking stagiair Kobe Martens – stage vanaf 16 november"),
        ("2", "", ""),
        ("3", "", ""),
    ]
    for i, (nr, mis, beter) in enumerate(rijen, start=1):
        cells = t.rows[i].cells
        cells[0].text = nr
        cells[1].text = mis
        cells[2].text = beter
        hoogte(t.rows[i], 1.1)
    breedtes(t, [0.8, 7.6, 8.6])
    doc.add_paragraph()

    # ---- Deel 2
    kop(doc, "Deel 2 — Aan, Cc of Bcc?")
    tekst(doc, "Schrijf bij elke situatie Aan, Cc of Bcc.", klein=True)
    t = doc.add_table(rows=4, cols=2)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(["Situatie", "Aan, Cc of Bcc?"]):
        c = t.rows[0].cells[i]
        c.text = h
        c.paragraphs[0].runs[0].bold = True
        schaduw(c)
    situaties = [
        "a.  Je mailt je mentor. Zij moet je vraag beantwoorden.",
        "b.  Je stagebegeleider van school wil je e-mail lezen, maar moet zelf niets doen.",
        "c.  De Speelboom stuurt één e-mail naar de ouders van alle 20 kinderen. De ouders mogen elkaars "
        "e-mailadres niet zien.",
    ]
    for i, s in enumerate(situaties, start=1):
        t.rows[i].cells[0].text = s
        hoogte(t.rows[i], 0.9)
    breedtes(t, [12.5, 4.5])
    doc.add_paragraph()
    tekst(doc, "Vraag 1. Waarom mogen de ouders in situatie c elkaars e-mailadres niet zien?", vet=True)
    antwoordlijnen(doc, 2)

    # ---- Deel 3
    kop(doc, "Deel 3 — Mijn e-mail aan mijn mentor")
    tekst(doc, "Klik in een vak en typ. De grijze woorden links zeggen wat in elk vak hoort. Laat Bcc leeg.",
          klein=True)
    venster(doc, ["aanspreking", "wie ben je?", "waarom mail je?", "je vraag", "slotgroet", "je naam,\nklas en school"])
    doc.add_paragraph()

    # ---- Deel 4
    kop(doc, "Deel 4 — Nalezen")
    tekst(doc, "Kruis aan wat klopt. Klopt iets nog niet? Verbeter je e-mail eerst.", klein=True)
    for s in [
        "Bij Aan staat het adres van mijn mentor.",
        "Bij Cc staat het adres van mijn stagebegeleider. Bcc is leeg.",
        "Mijn onderwerp zegt in een paar woorden waarover mijn e-mail gaat.",
        "Mijn aanspreking is: Beste mevrouw Verhoeven,",
        "Ik schrijf overal u en uw, niet je en jij.",
        "Ik zeg wie ik ben, waarom ik mail en wat ik vraag.",
        "Mijn slotgroet is: Met vriendelijke groeten, met daaronder mijn voornaam en achternaam, klas en school.",
        "Geen emoji's, geen chattaal, geen afkortingen zoals mvg.",
    ]:
        p = doc.add_paragraph("☐  " + s)
        p.paragraph_format.space_after = Pt(2)
    doc.add_paragraph()
    tekst(doc, "Je testlezer: laat je buur je e-mail lezen. Welke tip gaf je buur?", vet=True)
    antwoordlijnen(doc, 1)
    tekst(doc, "Wat heb je daardoor aangepast?", vet=True)
    antwoordlijnen(doc, 1)

    # ---- Deel 5
    kop(doc, "Deel 5 — Slotvraag")
    tekst(doc, "Vraag 2. Je mentor krijgt elke dag veel e-mails. Waarom helpt een duidelijk onderwerp haar?",
          vet=True)
    antwoordlijnen(doc, 2)

    doc.save(os.path.join(HERE, "Stagemail.docx"))


if __name__ == "__main__":
    stagemail()
    print("Gemaakt: Stagemail.docx")
